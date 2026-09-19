"""Handover formats: the intake pack, and the Word/PowerPoint deliverables."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from helpers import SAMPLE_CONFIG                                  # noqa: F401
from trailguide.config import Config
from trailguide.connectors import registered_connectors
from trailguide.intake import (
    ClientProfile, intake_sheets, read_workbook, render_config, render_readme,
    write_csv_stubs, write_workbook,
)
from trailguide.pipeline import run as run_pipeline
from trailguide.report.documents import _is_divider, _money, markdown_to_docx


class TestIntakeSpec(unittest.TestCase):
    """The template is generated from the connectors, so it cannot drift."""

    def test_every_sheet_declares_what_it_needs(self):
        sheets = intake_sheets()
        self.assertTrue(sheets)
        for sheet in sheets:
            with self.subTest(sheet=sheet.key):
                self.assertTrue(sheet.columns, "no columns declared")
                self.assertTrue(sheet.export_from, "no export instructions")
                self.assertIn(sheet.priority, ("core", "recommended", "optional"))
                self.assertTrue(any(c.required for c in sheet.columns),
                                "no column marked required")

    def test_sheet_keys_are_unique(self):
        keys = [sheet.key for sheet in intake_sheets()]
        self.assertEqual(len(keys), len(set(keys)))

    def test_every_declared_source_type_resolves_to_a_connector(self):
        from trailguide.connectors import get_connector
        for sheet in intake_sheets():
            with self.subTest(sheet=sheet.key):
                self.assertIsNotNone(get_connector(sheet.source_type))

    def test_the_core_sources_are_declared(self):
        """Search Console, GA4 and a crawl are what the analysis cannot run without."""
        core = {sheet.source_type for sheet in intake_sheets() if sheet.priority == "core"}
        self.assertLessEqual({"gsc", "ga4", "screaming_frog"}, core)

    def test_connectors_that_read_files_declare_intake(self):
        """A connector nobody can supply data for is a gap in the pack."""
        undeclared = [
            connector.name for connector in registered_connectors()
            if not connector.intake and connector.name != "generic"
        ]
        self.assertEqual(undeclared, [])


class TestIntakePack(unittest.TestCase):
    def setUp(self):
        self.profile = ClientProfile(
            name="Acme Co", domain="acme.com", brand_terms=("acme",),
            competitors=("rival.com",), revenue_target=250000,
        )

    def test_slug_is_filesystem_safe(self):
        self.assertEqual(ClientProfile(name="Acme & Co!").slug, "acme-co")
        self.assertEqual(ClientProfile(name="   ").slug, "client")

    def test_generated_config_loads_and_lists_every_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "acme.yml"
            path.write_text(render_config(self.profile), encoding="utf-8")
            config = Config.load(path)
            self.assertEqual(config.client_name, "Acme Co")
            self.assertEqual(config.domain, "acme.com")
            self.assertEqual(len(config.sources), len(intake_sheets()))

    def test_ecommerce_config_carries_ecommerce_economics(self):
        rendered = render_config(ClientProfile(name="Shop", economics_model="ecommerce"))
        self.assertIn("average_order_value", rendered)
        self.assertNotIn("average_contract_value", rendered)

    def test_csv_stubs_carry_the_declared_headers(self):
        import csv
        with tempfile.TemporaryDirectory() as tmp:
            written = write_csv_stubs(tmp)
            self.assertEqual(len(written), len(intake_sheets()))
            by_name = {path.name: path for path in written}
            for sheet in intake_sheets():
                with self.subTest(sheet=sheet.key):
                    with by_name[sheet.filename].open(encoding="utf-8") as handle:
                        self.assertEqual(next(csv.reader(handle)), sheet.headers)

    def test_readme_names_every_export(self):
        readme = render_readme(self.profile)
        for sheet in intake_sheets():
            self.assertIn(sheet.filename, readme)

    def test_workbook_round_trips_through_collect(self):
        """Fill a tab, extract it, and the CSV a connector reads comes back out."""
        import openpyxl
        sheets = intake_sheets()
        target = next(s for s in sheets if s.key == "search_console_queries")
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "intake.xlsx"
            write_workbook(book, self.profile, sheets)

            workbook = openpyxl.load_workbook(book)
            worksheet = workbook[target.title[:31].strip()]
            worksheet.append(["shoes", "https://acme.com/a", 10, 100, "10%", 4.2])
            workbook.save(book)

            written = read_workbook(book, Path(tmp) / "data")
            self.assertEqual([path.name for path in written], [target.filename])
            rows = written[0].read_text(encoding="utf-8").strip().splitlines()
            self.assertEqual(len(rows), 2)
            self.assertIn("shoes", rows[1])

    def test_empty_workbook_yields_no_csvs(self):
        """A blank pack is not an error - it just has nothing to extract yet."""
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "intake.xlsx"
            write_workbook(book, self.profile)
            self.assertEqual(read_workbook(book, Path(tmp) / "data"), [])

    def test_workbook_data_tabs_have_no_example_rows(self):
        """An example row survives into a real audit and quietly corrupts it."""
        import openpyxl
        with tempfile.TemporaryDirectory() as tmp:
            book = Path(tmp) / "intake.xlsx"
            write_workbook(book, self.profile)
            workbook = openpyxl.load_workbook(book)
            titles = {sheet.title[:31].strip() for sheet in intake_sheets()}
            for title in titles:
                self.assertEqual(workbook[title].max_row, 1, f"{title} has data rows")


class TestDocumentHelpers(unittest.TestCase):
    def test_money_is_compact(self):
        self.assertEqual(_money(1_297_500), "$1.30M")
        self.assertEqual(_money(142_786), "$143K")
        self.assertEqual(_money(420), "$420")
        self.assertEqual(_money(2_500_000_000), "$2.50B")

    def test_table_divider_detection(self):
        self.assertTrue(_is_divider("| --- | ---: |"))
        self.assertTrue(_is_divider("| :---: |"))
        self.assertFalse(_is_divider("| Metric | Value |"))

    def test_markdown_renders_headings_tables_and_lists(self):
        from docx import Document
        document = Document()
        markdown_to_docx(
            "# Title\n\n## Section\n\nA **bold** and _italic_ line.\n\n"
            "| Metric | Value |\n| --- | ---: |\n| Revenue | $1.2M |\n\n"
            "- first bullet\n- second bullet\n\n1. step one\n2. step two\n",
            document,
        )
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertEqual(headings, ["Title", "Section"])
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].rows[1].cells[1].text, "$1.2M")
        bullets = [p.text for p in document.paragraphs if p.style.name == "List Bullet"]
        self.assertEqual(bullets, ["first bullet", "second bullet"])
        numbered = [p.text for p in document.paragraphs if p.style.name == "List Number"]
        self.assertEqual(numbered, ["step one", "step two"])

    def test_inline_emphasis_becomes_runs(self):
        from docx import Document
        document = Document()
        markdown_to_docx("A **bold** word and `code`.", document)
        runs = document.paragraphs[0].runs
        self.assertTrue(any(run.bold and run.text == "bold" for run in runs))
        self.assertTrue(any(run.text == "code" for run in runs))


class TestDeliverables(unittest.TestCase):
    """End to end against the bundled sample data."""

    @classmethod
    def setUpClass(cls):
        cls.result = run_pipeline(Config.load(SAMPLE_CONFIG))

    def test_docx_carries_the_report_and_the_appendix(self):
        from docx import Document
        from trailguide.report import write_docx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_docx(self.result, Path(tmp) / "a.docx", max_detail=3)
            document = Document(path)
            headings = [
                p.text for p in document.paragraphs if p.style.name.startswith("Heading")
            ]
            self.assertTrue(any("Executive summary" in h for h in headings))
            self.assertTrue(any("Appendix" in h for h in headings))
            self.assertTrue(document.tables)

    def test_docx_can_omit_the_appendix(self):
        from docx import Document
        from trailguide.report import write_docx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_docx(self.result, Path(tmp) / "a.docx", 3, include_appendix=False)
            headings = [p.text for p in Document(path).paragraphs]
            self.assertFalse(any("Appendix" in h for h in headings))

    def test_pptx_covers_the_questions_asked_in_the_room(self):
        from pptx import Presentation
        from trailguide.report import write_pptx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx(self.result, Path(tmp) / "a.pptx")
            presentation = Presentation(path)
            self.assertGreaterEqual(len(presentation.slides), 7)
            text = " ".join(
                shape.text_frame.text
                for slide in presentation.slides
                for shape in slide.shapes
                if shape.has_text_frame
            )
            for expected in ("What is on the table", "The first moves",
                             "How we will know it worked", "What this does not claim"):
                self.assertIn(expected, text)

    def test_pptx_is_widescreen(self):
        from pptx import Presentation
        from trailguide.report import write_pptx
        with tempfile.TemporaryDirectory() as tmp:
            path = write_pptx(self.result, Path(tmp) / "a.pptx")
            presentation = Presentation(path)
            ratio = presentation.slide_width / presentation.slide_height
            self.assertAlmostEqual(ratio, 16 / 9, places=2)

    def test_jira_csv_has_one_row_per_scheduled_item(self):
        import csv
        from trailguide.report import write_jira_csv
        with tempfile.TemporaryDirectory() as tmp:
            path = write_jira_csv(self.result, Path(tmp) / "j.csv")
            with path.open(encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(len(rows), len(self.result.portfolio.scheduled))
            self.assertTrue(all(row["Summary"] and row["Issue Type"] for row in rows))


if __name__ == "__main__":
    unittest.main()
